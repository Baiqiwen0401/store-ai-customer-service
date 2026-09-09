from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("个体门店AI客服生产级技术方案_v2.0.docx")
FONT = "Microsoft YaHei"
BLACK = "000000"
NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "D9E1E8"
MUTED = "5B6573"
WHITE = "FFFFFF"


def font(run, size=10.5, bold=False, color=BLACK, italic=False, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row_index, row in enumerate(table.rows):
        if header and row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header_flag = OxmlElement("w:tblHeader")
            header_flag.set(qn("w:val"), "true")
            tr_pr.append(header_flag)
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    font(run, size=9.2, bold=header and row_index == 0, color=WHITE if header and row_index == 0 else BLACK)


def table(doc, headers, rows, widths):
    result = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        paragraph = result.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        font(run, size=9.2, bold=True, color=WHITE)
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(str(value))
            font(run, size=9.2)
    style_table(result, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return result


def paragraph(doc, text="", size=10.5, bold=False, color=BLACK, before=0, after=7, align=None, italic=False):
    result = doc.add_paragraph()
    if align is not None:
        result.alignment = align
    result.paragraph_format.space_before = Pt(before)
    result.paragraph_format.space_after = Pt(after)
    result.paragraph_format.line_spacing = 1.28
    if text:
        run = result.add_run(text)
        font(run, size=size, bold=bold, color=color, italic=italic)
    return result


def rich_paragraph(doc, parts, before=0, after=7):
    result = doc.add_paragraph()
    result.paragraph_format.space_before = Pt(before)
    result.paragraph_format.space_after = Pt(after)
    result.paragraph_format.line_spacing = 1.28
    for text, kwargs in parts:
        run = result.add_run(text)
        font(run, **kwargs)
    return result


def bullet(doc, text, level=0):
    result = doc.add_paragraph(style="List Bullet")
    result.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    result.paragraph_format.first_line_indent = Inches(-0.16)
    result.paragraph_format.space_after = Pt(4)
    result.paragraph_format.line_spacing = 1.2
    run = result.add_run(text)
    font(run, size=10.3)
    return result


def numbered(doc, text):
    result = doc.add_paragraph(style="List Number")
    result.paragraph_format.left_indent = Inches(0.25)
    result.paragraph_format.first_line_indent = Inches(-0.16)
    result.paragraph_format.space_after = Pt(4)
    result.paragraph_format.line_spacing = 1.2
    run = result.add_run(text)
    font(run, size=10.3)
    return result


def heading(doc, text, level=1):
    result = doc.add_paragraph(style=f"Heading {level}")
    result.paragraph_format.keep_with_next = True
    run = result.add_run(text)
    font(run, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLACK)
    return result


def flow(doc, text):
    result = doc.add_paragraph()
    result.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result.paragraph_format.space_before = Pt(3)
    result.paragraph_format.space_after = Pt(10)
    result.paragraph_format.line_spacing = 1.2
    run = result.add_run(text)
    font(run, size=10, bold=True, color=NAVY, name="Consolas")
    return result


def page_break(doc):
    doc.add_page_break()


def setup_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28
    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(17 if level == 1 else 11 if level == 2 else 7)
        style.paragraph_format.space_after = Pt(7 if level == 1 else 4)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.3)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("个体门店 AI 客服生产级技术方案  |  v2.0  |  2026-09-09")
    font(run, size=8, color=MUTED)


def build():
    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(55)
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run("个体门店 AI 客服生产级技术方案")
    font(title_run, size=24, bold=True, color=BLACK)
    paragraph(doc, "面向美容院网页客服的 AI 员工与人工协同系统", size=13, color=MUTED, after=26, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "生产基线版本 2.0", size=11, bold=True, color=NAVY, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "适用对象：单门店生产部署，可平滑扩展至多门店", size=10.5, color=MUTED, after=42, align=WD_ALIGN_PARAGRAPH.CENTER)
    table(doc, ["文档属性", "内容"], [
        ("方案状态", "推荐采用；用于生产级设计、开发和验收"),
        ("核心决策", "自有客服业务后端 + Dify Workflow/RAG + 生产基础设施"),
        ("首期渠道", "网页客服；后续可接入公众号、企业微信、小程序"),
        ("首期业务", "项目咨询、价格咨询、注意事项、预约意向、人工接管"),
        ("模型接入", "OpenAI-compatible 模型网关；密钥仅通过运行环境注入"),
        ("编制日期", "2026 年 9 月 9 日"),
    ], [1800, 7560])
    paragraph(doc, "本方案的核心结论：市场上成熟的门店 AI 客服普遍采用“客服工作台 + AI Agent + RAG 知识库 + 业务工具 + 人工兜底”的组合。系统不能把大模型当作唯一业务中枢，风险规则、客户数据、人工接管、预约任务和审计必须由自有业务后端控制。", size=11, after=8)
    paragraph(doc, "本文件同时定义了上线前必须完成的生产化改造、门店方需要提供的资料与决策，以及可量化的验收标准。", size=11, after=0)
    page_break(doc)

    heading(doc, "目录", 1)
    for item in [
        "1 方案摘要与决策边界",
        "2 目标、范围与非目标",
        "3 总体架构与生产部署",
        "4 客服请求处理流程",
        "5 AI 工作流与 RAG 设计",
        "6 长期记忆与持续专业化",
        "7 人工协同与预约意向",
        "8 数据模型与接口边界",
        "9 安全、隐私与合规",
        "10 可用性、性能与运维",
        "11 研发实施计划与验收",
        "12 你需要准备和确认的事项",
        "13 风险、取舍与后续演进",
        "附录 A 参考产品与技术选型依据",
        "附录 B 首期评测集建议",
    ]:
        paragraph(doc, item, size=10.5, after=4)
    page_break(doc)

    heading(doc, "1 方案摘要与决策边界", 1)
    paragraph(doc, "本项目建设一个面向美容院门店的网页 AI 客服。客户可在线询问项目、价格、营业时间、地址、护理注意事项并表达预约意向；系统优先使用门店已发布资料回答，复杂问题才进入 RAG 和大模型分析，涉及健康风险、投诉、退款或资料不足时转人工。门店员工在工作台中处理会话、任务和记忆审核。")
    heading(doc, "1.1 推荐方案", 2)
    table(doc, ["层次", "推荐技术", "职责"], [
        ("客户入口", "网页客服 Widget + HTTPS + SSE/WebSocket", "消息输入、流式展示、断线恢复、访客会话"),
        ("业务后端", "Python/FastAPI 或等价 API 服务", "租户、客户、会话、人工、预约任务、权限、审计"),
        ("AI 编排", "Dify Workflow API", "RAG 检索、模型调用、提示词、答案结构化输出"),
        ("模型网关", "OpenAI-compatible Adapter", "密钥隔离、超时、重试、熔断、备用模型"),
        ("知识检索", "混合检索 + 向量库/pgvector + 可选重排", "门店资料检索、版本、发布、回滚"),
        ("数据层", "PostgreSQL + Redis + 对象存储", "交易数据、缓存/队列、文档与附件"),
        ("部署层", "Docker Compose 起步，生产可迁移 Kubernetes", "反向代理、TLS、备份、日志、监控"),
    ], [1500, 2850, 5010])
    heading(doc, "1.2 必须坚持的边界", 2)
    bullet(doc, "风险规则优先于模型：过敏、红肿、破损、疾病、孕期、医美注射、退款、投诉和纠纷等场景不得由模型自由回答。")
    bullet(doc, "事实来源可追溯：项目、价格、营业时间、政策和地址必须来自已发布知识或业务系统，不允许模型补写未配置事实。")
    bullet(doc, "AI 不直接确认正式预约：首期只登记预约意向，门店确认后才形成正式预约。")
    bullet(doc, "长期记忆必须经过同意和审核：模型提取的是候选事实，不是自动写入的永久用户画像。")
    bullet(doc, "Dify 是 AI 能力服务，不替代客服业务后端、权限系统和人工工作台。")

    heading(doc, "2 目标、范围与非目标", 1)
    heading(doc, "2.1 首期目标", 2)
    table(doc, ["目标", "生产要求", "首期结果"], [
        ("自动回答", "可追溯、可降级、可观察", "项目、价格、地址、营业时间、注意事项等标准问题自动回答"),
        ("意向转化", "字段校验、重复去重、任务可跟进", "收集称呼、手机号、日期/时段、项目偏好并形成待办"),
        ("人工协作", "上下文完整、状态明确、可审计", "员工认领、回复、暂停 AI、恢复 AI、关闭会话"),
        ("知识运营", "草稿、审核、发布、版本、回滚", "员工可维护门店 FAQ 和服务资料"),
        ("持续专业化", "可控的记忆和反馈闭环", "客户偏好候选审核后用于后续回答；低质量答案可标记"),
    ], [1800, 3600, 3960])
    heading(doc, "2.2 明确不在首期范围", 2)
    bullet(doc, "不做疾病诊断、疗效保证、医美注射建议或治疗方案。")
    bullet(doc, "不做自动扣款、退款、正式排班锁定或其他不可逆业务动作。")
    bullet(doc, "不在首期同时接入小红书、微信、抖音等平台；渠道接入需使用平台官方能力并单独验收。")
    bullet(doc, "不使用线上自动微调替代知识维护；首期以知识版本、人工反馈和记忆审核实现持续改进。")

    heading(doc, "3 总体架构与生产部署", 1)
    flow(doc, "网页 Widget -> API Gateway -> 会话服务 -> 风险/意图路由 -> 直答或 Dify Workflow -> 人工/任务系统")
    heading(doc, "3.1 逻辑架构", 2)
    table(doc, ["组件", "职责", "关键生产要求"], [
        ("API Gateway", "TLS、限流、请求体大小、来源校验", "Nginx/云负载均衡；禁止直接暴露应用端口"),
        ("Conversation Service", "消息、会话状态、幂等、SSE 推送", "消息 ID、幂等键、顺序号、连接恢复"),
        ("Policy Router", "风险、意图、答案可支撑性门禁", "规则可配置、版本化、单元测试"),
        ("AI Orchestrator", "调用 Dify、解析结构化输出、降级", "超时、重试、熔断、trace_id、模型路由"),
        ("Knowledge Service", "知识导入、切分、索引、发布、回滚", "草稿/发布分离、版本和来源追踪"),
        ("Human Workspace", "员工会话、任务、标签、记忆审核", "RBAC、审计、操作防重复"),
        ("Data Services", "PostgreSQL、Redis、对象存储、向量检索", "备份、恢复演练、加密、容量监控"),
        ("Observability", "日志、指标、链路、评测", "脱敏、告警、按租户查询、保留策略"),
    ], [1800, 3600, 3960])
    heading(doc, "3.2 推荐生产部署", 2)
    flow(doc, "HTTPS/Nginx -> API 服务 x2 -> PostgreSQL 主库 + 备份 -> Redis -> Worker -> Dify -> 模型中转站")
    paragraph(doc, "单门店可先使用一台具备持久化磁盘的生产主机运行 Docker Compose，但必须使用 PostgreSQL、Redis、备份、反向代理和监控。应用服务应保持无状态，未来可水平扩展。Dify 可独立部署或使用受控的内部实例；模型密钥只注入服务运行环境，不进入前端、数据库或 Git。")
    heading(doc, "3.3 环境分层", 2)
    table(doc, ["环境", "用途", "数据规则"], [
        ("开发", "开发联调和自动化测试", "使用脱敏或模拟数据，禁止真实客户信息"),
        ("预发布", "生产配置等价验证、验收和压测", "使用匿名化评测集，模型调用可设置预算"),
        ("生产", "真实客户服务", "独立密钥、独立数据库、备份、监控和审计"),
    ], [1500, 3600, 4260])

    heading(doc, "4 客服请求处理流程", 1)
    flow(doc, "接收消息 -> 幂等校验 -> 风险检测 -> 意图识别 -> 结构化直答 -> RAG 答案门禁 -> 模型回答 -> 任务/记忆/审计")
    heading(doc, "4.1 路由优先级", 2)
    table(doc, ["优先级", "判断", "处理"], [
        ("P0", "会话已由人工接管或已关闭", "不调用 AI，提示人工状态"),
        ("P1", "健康/投诉/退款/纠纷等风险", "安全话术 + 转人工 + 创建跟进任务"),
        ("P2", "价格、项目、地址、营业时间等高置信意图", "读取已发布结构化知识直答"),
        ("P3", "预约意向", "收集必要字段并创建/更新待办，不直接承诺成功"),
        ("P4", "资料覆盖且需要自然语言组织", "进入 Dify RAG Workflow"),
        ("P5", "资料不足或模型不可用", "明确说明资料范围，创建人工确认任务"),
    ], [1200, 3900, 4260])
    heading(doc, "4.2 标准回复原则", 2)
    bullet(doc, "资料命中时先回答客户问题，再补充下一步行动，不先输出系统故障。")
    bullet(doc, "模型不可用时只影响复杂咨询，不影响项目、价格、地址、营业时间等结构化直答。")
    bullet(doc, "涉及治疗、疾病或风险时统一使用安全话术：AI 回复不作为治疗依据，建议转人工评估。")
    bullet(doc, "每次 AI 回复记录 intent、route、knowledge_version、confidence、model_latency 和 trace_id。")

    heading(doc, "5 AI 工作流与 RAG 设计", 1)
    heading(doc, "5.1 Dify Workflow 节点", 2)
    table(doc, ["节点", "输入/输出", "控制要求"], [
        ("Start", "message、tenant_id、conversation_id、customer_context", "校验字段和长度，拒绝非法租户"),
        ("Context", "门店资料、已审核记忆、会话摘要", "只注入当前租户和已授权数据"),
        ("Knowledge Retrieval", "问题 -> 相关片段", "混合检索；返回来源、版本、分数"),
        ("Answerability Gate", "片段相关性 -> can_answer", "不足时禁止模型编造"),
        ("LLM Response", "问题 + 证据 -> structured answer", "输出 answer、handoff、citations、confidence"),
        ("Safety Check", "模型结果 -> safe/blocked", "检测医疗承诺、越权动作和敏感信息"),
        ("End", "统一 JSON 结果", "由自有后端落库、推送、建任务"),
    ], [2100, 3300, 3960])
    heading(doc, "5.2 知识库分层", 2)
    table(doc, ["知识类型", "示例", "维护方式"], [
        ("门店事实", "项目、价格、地址、营业时间、联系方式", "结构化字段 + FAQ，发布后生效"),
        ("服务说明", "护理流程、适用人群、注意事项", "文档/FAQ，必须标注来源和更新时间"),
        ("业务政策", "预约、取消、退款、投诉处理", "版本化政策，变更需员工审核"),
        ("安全边界", "禁忌、风险词、转人工话术", "规则配置，变更需测试和审批"),
        ("客户记忆", "肤质、预算、时间偏好", "单独存储，客户同意并经员工审核"),
    ], [1800, 3960, 3600])
    paragraph(doc, "RAG 采用混合检索：关键词/全文检索用于项目名、价格和政策编号，向量检索用于自然语言表达，重排模型用于提升候选片段排序。最终由答案可支撑性门禁决定是否允许模型回答。知识库应支持草稿、审核、发布、回滚和生效时间。")
    heading(doc, "5.3 结构化输出契约", 2)
    paragraph(doc, "Dify Workflow 不返回不可解析的自由文本作为系统控制信号，必须返回 JSON。建议字段为：answer、route、intent、confidence、can_answer、handoff_required、handoff_reason、citations、suggested_task、memory_candidates、model_usage、trace_id。业务后端只信任通过 schema 校验和安全规则的字段。")

    heading(doc, "6 长期记忆与持续专业化", 1)
    flow(doc, "客户表达 -> 候选事实提取 -> 客户同意校验 -> 员工审核 -> approved memory -> 后续回答上下文")
    table(doc, ["记忆层", "内容", "生命周期"], [
        ("会话记忆", "当前对话上下文、摘要、未完成任务", "会话关闭或按策略归档"),
        ("客户偏好", "肤质、预算、时间偏好、关注项目", "候选 -> 审核 -> 生效 -> 可撤回"),
        ("门店知识", "项目、价格、政策、服务标准", "草稿 -> 审核 -> 发布 -> 版本回滚"),
        ("运营反馈", "人工修改、转人工原因、低置信问题", "进入评测集和知识改进队列"),
    ], [1800, 3960, 3600])
    bullet(doc, "客户未授权时，不保存手机号和长期偏好；会话临时上下文仍需遵循数据保留策略。")
    bullet(doc, "模型只能提出 memory candidate，不能绕过后端直接写入 approved memory。")
    bullet(doc, "员工可查看来源消息、置信度、创建时间和审核人，并能撤回已批准记忆。")
    bullet(doc, "“越用越专业”主要通过知识迭代、反馈评测和审核记忆实现，不承诺模型自动训练。")

    heading(doc, "7 人工协同与预约意向", 1)
    heading(doc, "7.1 人工接管状态机", 2)
    flow(doc, "open -> handoff_pending -> human_active -> ai_paused -> ai_resumed -> closed")
    table(doc, ["状态", "AI 行为", "员工行为"], [
        ("open", "按路由正常回答", "可查看会话和任务"),
        ("handoff_pending", "停止继续生成，发送等待提示", "认领并查看摘要、来源和风险原因"),
        ("human_active", "不自动回复客户", "人工回复、添加标签、创建任务"),
        ("ai_resumed", "重新按规则处理", "确认人工问题已解决"),
        ("closed", "不再处理原会话消息", "可由客户重新发起新会话"),
    ], [1800, 3600, 3960])
    heading(doc, "7.2 预约意向字段", 2)
    table(doc, ["字段", "规则", "是否必填"], [
        ("客户称呼", "文本长度 1–30，禁止敏感信息", "建议"),
        ("手机号", "脱敏展示，格式校验，取得同意后保存", "转化跟进必填"),
        ("意向项目", "仅允许已发布项目或人工备注", "建议"),
        ("期望日期/时段", "自然语言解析后由员工确认", "预约意向必填"),
        ("备注", "限制长度，自动脱敏", "可选"),
    ], [2100, 4500, 2760])

    heading(doc, "8 数据模型与接口边界", 1)
    heading(doc, "8.1 核心实体", 2)
    table(doc, ["实体", "关键字段", "归属"], [
        ("Tenant/Store", "tenant_id、名称、营业时间、配置版本", "门店"),
        ("Customer", "customer_id、手机号哈希/加密值、同意状态", "门店"),
        ("Conversation", "状态、渠道、AI 开关、assigned_to、trace_id", "门店"),
        ("Message", "role、content、confidence、route、metadata", "会话"),
        ("KnowledgeDocument", "版本、来源、状态、发布时间、checksum", "门店"),
        ("Memory", "类型、内容、来源、状态、审核人、撤回时间", "客户"),
        ("Task", "类型、摘要、负责人、状态、截止时间", "门店"),
        ("AuditLog", "actor、action、entity、before/after、created_at", "门店/平台"),
    ], [1800, 4860, 2700])
    heading(doc, "8.2 内部 API 原则", 2)
    bullet(doc, "所有写接口支持 Idempotency-Key；重复请求不得重复创建消息、任务或预约意向。")
    bullet(doc, "所有租户数据查询必须由服务端从认证上下文取得 tenant_id，不信任前端传入值。")
    bullet(doc, "AI 服务通过内部 API 或签名请求调用；Dify 不直接访问数据库。")
    bullet(doc, "模型失败返回可分类错误：timeout、auth、rate_limit、network、invalid_response、budget_exceeded。")
    bullet(doc, "异步任务采用队列，任务状态可重试、可取消、可观察，不能依靠单次 HTTP 请求完成所有工作。")

    heading(doc, "9 安全、隐私与合规", 1)
    table(doc, ["领域", "生产要求", "验收证据"], [
        ("身份与权限", "员工登录、RBAC、最小权限、会话超时", "权限矩阵和越权测试"),
        ("密钥管理", "密钥使用环境变量或密钥服务，禁止前端暴露", "仓库扫描、运行时配置检查"),
        ("隐私保护", "手机号加密/脱敏、同意记录、删除和导出", "数据处理流程和删除演练"),
        ("租户隔离", "所有查询、缓存、向量检索带 tenant_id", "跨租户访问测试"),
        ("输入安全", "长度限制、提示注入防护、文件类型和大小校验", "恶意输入测试"),
        ("审计", "知识发布、人工操作、记忆审核、工具调用全记录", "审计日志抽查"),
        ("内容安全", "医疗承诺、歧视、隐私泄露和越权动作拦截", "安全评测集"),
        ("数据留存", "定义消息、日志、备份和已撤回记忆的保留期限", "配置和定期清理记录"),
    ], [1700, 4560, 3100])
    paragraph(doc, "美容院涉及皮肤健康和个人信息，系统必须在产品文案、隐私政策和人工服务流程中明确：AI 回复不作为治疗依据，涉及疾病、过敏或医美项目时建议转人工评估。具体法律义务需结合实际经营主体、地区和渠道规则由专业人员确认。")

    heading(doc, "10 可用性、性能与运维", 1)
    heading(doc, "10.1 建议 SLO", 2)
    table(doc, ["指标", "首期目标", "监控方式"], [
        ("结构化问题 P95 首字节", "≤ 1.5 秒", "API latency histogram"),
        ("RAG/模型问题 P95 完成时间", "≤ 8 秒，超时进入降级", "trace + model_events"),
        ("人工接管成功率", "≥ 99% 的转接请求形成任务", "会话/任务对账"),
        ("服务可用性", "月度 ≥ 99.5%（单机起步）", "健康检查和告警"),
        ("知识直答准确率", "评测集 ≥ 95%", "发布前离线评测"),
        ("无依据回答率", "≤ 1%", "人工抽检 + 安全集"),
        ("备份恢复目标", "RPO ≤ 24h，RTO ≤ 4h", "恢复演练"),
    ], [2300, 2500, 4560])
    heading(doc, "10.2 生产运维最低项", 2)
    bullet(doc, "健康检查：API、数据库、Redis、Worker、Dify 和模型网关分别检查，不以进程存活代替业务可用。")
    bullet(doc, "日志：结构化 JSON 日志，包含 request_id、trace_id、tenant_id、conversation_id，并对手机号和密钥脱敏。")
    bullet(doc, "告警：错误率、模型超时、队列积压、数据库空间、备份失败、知识索引失败。")
    bullet(doc, "发布：数据库迁移可回滚，Dify Workflow 和知识库版本可回退，生产发布保留变更记录。")
    bullet(doc, "备份：数据库每日全量、必要时增量；备份必须异地或独立存储，并至少每月恢复演练一次。")

    heading(doc, "11 研发实施计划与验收", 1)
    table(doc, ["阶段", "主要工作", "完成标准"], [
        ("阶段 0 需求冻结", "门店资料、风险边界、预约字段、权限角色", "业务规则和资料清单签字确认"),
        ("阶段 1 生产骨架", "PostgreSQL、Redis、鉴权、配置、日志、备份", "三环境可部署，基础安全检查通过"),
        ("阶段 2 AI 接入", "Dify Workflow、RAG、答案门禁、模型网关", "结构化 JSON、超时降级、调用可追踪"),
        ("阶段 3 业务闭环", "人工工作台、预约任务、记忆审核、知识发布", "关键状态机和权限测试通过"),
        ("阶段 4 质量验收", "真实评测集、压测、安全测试、恢复演练", "达到 SLO 和上线门槛"),
        ("阶段 5 灰度上线", "单门店小流量、人工旁路、问题复盘", "连续观察期无阻断问题后扩大流量"),
    ], [1800, 4500, 3060])
    heading(doc, "11.1 上线门槛", 2)
    bullet(doc, "P0/P1 风险问题为零；未授权访问、跨租户读取、密钥泄露为零。")
    bullet(doc, "标准知识评测集准确率达到目标；资料不足问题不会伪造门店事实。")
    bullet(doc, "模型服务不可用时，结构化直答、人工接管和任务创建仍可工作。")
    bullet(doc, "备份恢复、知识回滚、人工接管和 AI 暂停/恢复完成演练。")
    bullet(doc, "门店员工完成工作台培训，并确认 AI 回复和人工话术边界。")

    heading(doc, "12 你需要准备和确认的事项", 1)
    paragraph(doc, "以下事项是项目能否达到生产标准的前置条件。技术实现可以由开发完成，但资料真实性、业务规则和上线责任必须由门店方确认。")
    heading(doc, "12.1 你需要提供的资料", 2)
    table(doc, ["资料", "最低内容", "用途"], [
        ("门店基础信息", "名称、地址、电话、营业时间、节假日规则", "结构化直答和页面展示"),
        ("服务项目", "项目名称、时长、价格、适用说明、禁忌和注意事项", "知识库和项目直答"),
        ("预约规则", "可预约时段、提前多久、改期/取消规则", "预约意向待办和人工确认"),
        ("服务政策", "退款、投诉、优惠、会员、隐私政策", "政策问答和转人工"),
        ("安全边界", "必须转人工的关键词和场景", "风险路由与安全话术"),
        ("人工信息", "员工账号、角色、接待时间、接管责任人", "权限和任务分配"),
        ("品牌内容", "欢迎语、语气、禁用词、称呼偏好", "Prompt 和网页客服体验"),
    ], [1800, 4500, 3060])
    heading(doc, "12.2 你需要确认的决策", 2)
    table(doc, ["决策项", "需要确认的选项", "默认建议"], [
        ("部署方式", "本地服务器 / 云服务器 / 托管 SaaS", "云服务器或门店自有服务器，生产数据可控"),
        ("Dify 方式", "自部署 / 受控云服务", "生产优先自部署或签署数据处理约定"),
        ("模型预算", "月度预算、单次超时、备用模型", "先设月度预算和 8 秒业务超时"),
        ("客户数据", "保存期限、手机号是否长期保存", "最小化保存，获得同意后保存"),
        ("人工规则", "何时转人工、谁负责、多久响应", "风险和资料不足自动转人工"),
        ("正式预约", "是否接入排班/预约系统", "首期只做预约意向，不自动锁时段"),
    ], [1800, 4500, 3060])
    heading(doc, "12.3 你需要参与的验收", 2)
    numbered(doc, "逐条确认项目、价格和政策资料，确保知识库事实准确。")
    numbered(doc, "提供至少 30–50 条真实客户常问问题和标准答案或处理方式。")
    numbered(doc, "确认哪些问题必须人工处理，并审核安全话术。")
    numbered(doc, "用真实员工角色测试认领、回复、暂停 AI、恢复 AI 和任务完成。")
    numbered(doc, "参加上线前灰度观察和每周知识/问题复盘。")

    heading(doc, "13 风险、取舍与后续演进", 1)
    table(doc, ["风险", "影响", "控制措施"], [
        ("资料过期或价格变更", "回答错误、客诉", "知识发布审核、版本、生效时间、回滚"),
        ("模型延迟或不可用", "复杂问题体验下降", "超时、熔断、结构化直答、人工降级"),
        ("模型生成医疗承诺", "安全和合规风险", "风险规则、答案门禁、安全检测、人工评估"),
        ("平台渠道限制", "接入不稳定或封禁", "优先官方 API，渠道适配层隔离"),
        ("Dify 许可证限制", "商业化或多租户受限", "按实际部署方式核验许可证和商业条款"),
        ("客户隐私泄露", "合规和声誉风险", "最小化采集、加密、脱敏、权限、审计"),
        ("业务动作误执行", "预约/退款等损失", "工具白名单、参数校验、确认和幂等"),
    ], [1900, 3000, 4460])
    heading(doc, "13.1 后续演进顺序", 2)
    paragraph(doc, "第一阶段完成网页客服和单店生产基线；第二阶段接入正式预约系统和更多知识来源；第三阶段再评估公众号、企业微信、小红书等渠道；第四阶段在租户隔离、计费、运营后台和许可证合规成熟后，才考虑对外提供多门店 SaaS。")
    paragraph(doc, "不建议在第一阶段同时解决多渠道、多门店、自动支付和复杂 CRM。扩大范围会显著增加数据隔离、平台审核、运维和合规成本。")

    heading(doc, "附录 A 参考产品与技术选型依据", 1)
    table(doc, ["项目/产品", "可借鉴能力", "适用判断"], [
        ("Dify", "Workflow、RAG、模型接入、API、可观测性", "作为 AI 编排服务；核验自定义许可证"),
        ("FastGPT", "中文知识库、RAG、可视化工作流", "可作为 Dify 的备选或对比 PoC"),
        ("MaxKB", "中文企业知识库、Agent、快速部署", "适合知识问答；GPLv3 需评估商用边界"),
        ("RAGFlow", "复杂文档解析和高质量 RAG", "知识底座参考，不作为首期客服工作台"),
        ("Chatwoot", "人工坐席、会话、全渠道工作台", "未来人工工作台或多渠道参考"),
        ("TGO/Agent Desk/ChatterMate", "AI 客服、RAG、人工转接、网页 Widget", "参考成熟产品结构，不直接替代业务后端"),
    ], [1800, 3960, 3600])
    paragraph(doc, "参考链接：Dify https://github.com/langgenius/dify；FastGPT https://github.com/labring/FastGPT；MaxKB https://github.com/1Panel-dev/MaxKB；RAGFlow https://github.com/infiniflow/ragflow；Chatwoot https://github.com/chatwoot/chatwoot；TGO https://github.com/tgoai/tgo。链接和许可证应在实际部署前再次核验。", size=9.5, color=MUTED)

    heading(doc, "附录 B 首期评测集建议", 1)
    table(doc, ["类别", "建议数量", "示例", "通过标准"], [
        ("项目与价格", "10", "有哪些项目、补水多少钱", "事实准确，引用已发布资料"),
        ("门店信息", "5", "地址、电话、营业时间", "字段准确，不调用模型也能回答"),
        ("预约意向", "8", "周六下午想做补水", "收集字段、建待办、不承诺成功"),
        ("注意事项", "8", "敏感肌能不能做", "保守回答，必要时转人工"),
        ("风险问题", "8", "过敏、红肿、注射、退款", "风险话术和人工任务正确"),
        ("资料外问题", "6", "未配置的项目或疗效", "不编造，说明范围并转人工"),
        ("对抗性输入", "5", "要求泄露提示词/手机号", "拒绝越权，数据不泄露"),
    ], [1800, 1200, 3900, 2460])
    paragraph(doc, "建议每次知识发布、Prompt 变更或模型更换后自动运行这组评测，并保存版本、得分、失败样例和人工处置结果。", size=10.5, after=0)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
